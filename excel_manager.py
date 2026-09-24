#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Excel import/export for devices, blacklist, manual radios and time-range backups."""
from __future__ import annotations
import io, ipaddress, json, sqlite3
from datetime import datetime, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEVICE_HEADERS=["IP","City","Vendor","Type","Hostname","Model","RouterOS Version","Firmware","MAC","SSID","Frequency","Channel","Bandwidth","Mode","Radio","TX Power","RX Power","Signal","Peer TX Signal","TX Rate","RX Rate","CCQ","TX CCQ","RX CCQ","SNR","Noise Floor","Temperature","CPU Load","Memory Usage","Uptime","SNMP","API","SSH","Last Seen","Status"]
RADIO_HEADERS=["IP","Name","Model","Vendor","Serial","SSID","Frequency","Channel","Bandwidth","Mode","Radio Name","TX Power","RX Power","Signal","Peer TX Signal","TX Rate","RX Rate","CCQ","SNR","Noise Floor","Notes"]
BLACK_HEADERS=["Entry","Description","Enabled"]

def text(v): return "" if v is None else str(v)
def sheet_style(ws):
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    fill=PatternFill("solid",fgColor="1F4E78")
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF");c.fill=fill;c.alignment=Alignment(horizontal="center",vertical="center")
    for col in range(1,ws.max_column+1):
        max_len=min(42,max(12,max((len(text(r[col-1].value)) for r in ws.iter_rows(min_row=1,max_row=min(ws.max_row,120))),default=12)+2))
        ws.column_dimensions[get_column_letter(col)].width=max_len

def add_sheet(wb,title,headers,rows):
    ws=wb.create_sheet(title);ws.append(headers)
    for row in rows:ws.append([row.get(h,"") if isinstance(row,dict) else row[i] for i,h in enumerate(headers)])
    sheet_style(ws);return ws

def json_load(v):
    if not v:return {}
    if isinstance(v,(dict,list)):return v
    try:return json.loads(v)
    except:return {}

def export_excel(db_path, hours=None, mode="full"):
    conn=sqlite3.connect(db_path,timeout=60);conn.row_factory=sqlite3.Row
    try:
        wb=Workbook();wb.remove(wb.active)
        if mode=="blacklist":
            rows=[dict(r) for r in conn.execute("SELECT entry,description,enabled FROM blacklist ORDER BY entry")];add_sheet(wb,"Blacklist",BLACK_HEADERS,rows)
        elif mode=="radios":
            rows=[dict(r) for r in conn.execute("SELECT ip_address as IP,name as Name,model as Model,vendor as Vendor,serial_number as Serial,ssid as SSID,frequency as Frequency,channel as Channel,bandwidth as Bandwidth,mode as Mode,radio_name as 'Radio Name',tx_power as 'TX Power',rx_power as 'RX Power',signal_strength as Signal,peer_tx_signal as 'Peer TX Signal',tx_rate as 'TX Rate',rx_rate as 'RX Rate',ccq as CCQ,snr as SNR,noise_floor as 'Noise Floor',notes as Notes FROM manual_radios ORDER BY ip_address")];add_sheet(wb,"Manual Radios",RADIO_HEADERS,rows)
        else:
            device_rows=[dict(r) for r in conn.execute("SELECT * FROM devices ORDER BY city_id,ip_address")]
            if hours is not None and hours>0:
                cutoff=(datetime.now()-timedelta(hours=hours)).strftime("%Y-%m-%d %H:%M:%S")
            else: cutoff=None
            mapped=[]
            for d in device_rows:
                mapped.append({"IP":d.get("ip_address"),"City":d.get("city"),"Vendor":d.get("vendor"),"Type":d.get("device_type"),"Hostname":d.get("hostname"),"Model":d.get("model"),"RouterOS Version":d.get("routeros_version"),"Firmware":d.get("firmware_version"),"MAC":d.get("mac_address"),"SSID":d.get("ssid"),"Frequency":d.get("frequency"),"Channel":d.get("channel"),"Bandwidth":d.get("bandwidth"),"Mode":d.get("mode"),"Radio":d.get("radio_name"),"TX Power":d.get("tx_power"),"RX Power":d.get("rx_power"),"Signal":d.get("signal_strength"),"Peer TX Signal":d.get("peer_tx_signal"),"TX Rate":d.get("tx_rate"),"RX Rate":d.get("rx_rate"),"CCQ":d.get("ccq"),"TX CCQ":d.get("tx_ccq"),"RX CCQ":d.get("rx_ccq"),"SNR":d.get("snr"),"Noise Floor":d.get("noise_floor"),"Temperature":d.get("temperature"),"CPU Load":d.get("cpu_load"),"Memory Usage":d.get("memory_usage"),"Uptime":d.get("uptime"),"SNMP":d.get("snmp_status"),"API":d.get("api_status"),"SSH":d.get("ssh_status"),"Last Seen":d.get("last_seen"),"Status":d.get("scan_status")})
            add_sheet(wb,"Devices",DEVICE_HEADERS,mapped)
            cities=[dict(r) for r in conn.execute("SELECT * FROM cities ORDER BY name")];add_sheet(wb,"Cities",list(cities[0].keys()) if cities else ["id","name","cidr","description","enabled","ssh_port","snmp_port","api_port","api_enabled","api_ssl_enabled","api_timeout","snmp_timeout","snmp_retries","ssh_timeout","ssh_credentials_json","snmp_communities_json","created_at","updated_at"],cities)
            scans_q="SELECT * FROM scan_logs"+(" WHERE scan_time>=?" if cutoff else "")+" ORDER BY id DESC";scans=[dict(r) for r in conn.execute(scans_q,(cutoff,) if cutoff else ()).fetchall()];add_sheet(wb,"Scan History",list(scans[0].keys()) if scans else ["id","scan_time","started_by","selection_json","total_ips_scanned","devices_found","devices_success","devices_failed","devices_skipped","duration_seconds"],scans)
            ids=[r["id"] for r in scans]
            if ids:
                q=','.join('?'*len(ids));targets=[dict(r) for r in conn.execute(f"SELECT * FROM scan_results WHERE scan_id IN ({q}) ORDER BY scan_id DESC,city_name,ip_address",ids)];changes=[dict(r) for r in conn.execute(f"SELECT * FROM device_changes WHERE scan_id IN ({q}) ORDER BY changed_at DESC",ids)]
            else:targets=[];changes=[]
            add_sheet(wb,"Scan Targets",list(targets[0].keys()) if targets else ["id","scan_id","ip_address","device_id","city_id","city_name","status","stage","reason","snapshot_json","scanned_at"],targets)
            add_sheet(wb,"Changes",list(changes[0].keys()) if changes else ["id","scan_id","ip_address","field_name","old_value","new_value","changed_at"],changes)
            bl=[dict(r) for r in conn.execute("SELECT * FROM blacklist ORDER BY entry")];add_sheet(wb,"Blacklist",["id","entry","description","enabled","created_at"],bl)
            mr=[dict(r) for r in conn.execute("SELECT * FROM manual_radios ORDER BY ip_address")];add_sheet(wb,"Manual Radios",list(mr[0].keys()) if mr else RADIO_HEADERS,mr)
            mi=[dict(r) for r in conn.execute("SELECT * FROM manual_ips ORDER BY ip_address")];add_sheet(wb,"Manual IPs",list(mi[0].keys()) if mi else ["id","ip_address","description","enabled","created_at"],mi)
            if ids:
                latest=ids[0];rows=[dict(r) for r in conn.execute("SELECT ip_address,snapshot_json,status,city_name FROM scan_results WHERE scan_id=? ORDER BY city_name,ip_address",(latest,))]
                wr=[];reg=[];interfaces=[];queues=[];fw=[]
                for r in rows:
                    snap=json_load(r["snapshot_json"])
                    wr.append({"IP":r["ip_address"],"City":r["city_name"],"SSID":snap.get("ssid"),"Frequency":snap.get("frequency"),"Channel":snap.get("channel"),"Bandwidth":snap.get("bandwidth"),"Mode":snap.get("mode"),"Radio":snap.get("radio_name"),"TX Power":snap.get("tx_power"),"RX Power":snap.get("rx_power"),"Signal":snap.get("signal_strength"),"Peer TX Signal":snap.get("peer_tx_signal"),"TX Rate":snap.get("tx_rate"),"RX Rate":snap.get("rx_rate"),"CCQ":snap.get("ccq"),"TX CCQ":snap.get("tx_ccq"),"RX CCQ":snap.get("rx_ccq"),"SNR":snap.get("snr"),"Noise Floor":snap.get("noise_floor")})
                    for key,target,title in (("wireless_registration",reg,"Registration"),("interfaces",interfaces,"Interfaces"),("queues",queues,"Queues"),("firewall_counters",fw,"Firewall")):
                        if snap.get(key):target.append({"IP":r["ip_address"],"City":r["city_name"],title:snap.get(key) if isinstance(snap.get(key),str) else json.dumps(snap.get(key),ensure_ascii=False)})
                add_sheet(wb,"Wireless",list(wr[0].keys()) if wr else ["IP","City","SSID","Frequency","Channel","Bandwidth","Mode","Radio","TX Power","RX Power","Signal","Peer TX Signal","TX Rate","RX Rate","CCQ","TX CCQ","RX CCQ","SNR","Noise Floor"],wr)
                add_sheet(wb,"Registration",["IP","City","Registration"],reg);add_sheet(wb,"Interfaces",["IP","City","Interfaces"],interfaces);add_sheet(wb,"Queues",["IP","City","Queues"],queues);add_sheet(wb,"Firewall",["IP","City","Firewall"],fw)
        if not wb.worksheets: add_sheet(wb,"Info",["Message"],[{"Message":"No data"}])
        out=io.BytesIO();wb.save(out);wb.close();out.seek(0);return out
    finally:conn.close()

def import_excel(fileobj,db_path,mode="devices"):
    wb=load_workbook(fileobj,read_only=True,data_only=True)
    try:
        ws=wb.active;headers=[text(c.value).strip().lower() for c in next(ws.iter_rows(max_row=1))];pos={h:i for i,h in enumerate(headers)}
        conn=sqlite3.connect(db_path,timeout=60);conn.execute("BEGIN")
        try:
            if mode=="blacklist":
                p={"entry":pos.get("entry",pos.get("ip",pos.get("ip address"))),"description":pos.get("description"),"enabled":pos.get("enabled")};count=0
                if p["entry"] is None:raise ValueError("ستون Entry در Excel وجود ندارد")
                for row in ws.iter_rows(min_row=2,values_only=True):
                    entry=text(row[p["entry"]]).strip()
                    if not entry:continue
                    ipaddress.ip_network(entry,strict=False);desc=text(row[p["description"]]) if p["description"] is not None else "";enabled=int(row[p["enabled"]]) if p["enabled"] is not None and row[p["enabled"]] not in (None,"") else 1
                    conn.execute("INSERT INTO blacklist(entry,description,enabled,created_at) VALUES(?,?,?,datetime('now')) ON CONFLICT(entry) DO UPDATE SET description=excluded.description,enabled=excluded.enabled",(entry,desc,enabled));count+=1
                conn.commit();return {"status":"ok","updated":count}
            if mode=="radios":
                aliases={"ip":"ip_address","ip address":"ip_address","name":"name","model":"model","vendor":"vendor","serial":"serial_number","ssid":"ssid","frequency":"frequency","channel":"channel","bandwidth":"bandwidth","mode":"mode","radio name":"radio_name","tx power":"tx_power","rx power":"rx_power","signal":"signal_strength","peer tx signal":"peer_tx_signal","tx rate":"tx_rate","rx rate":"rx_rate","ccq":"ccq","snr":"snr","noise floor":"noise_floor","notes":"notes"};mp={aliases.get(h,h):i for h,i in pos.items()};count=0
                if "ip_address" not in mp:raise ValueError("ستون IP در Excel وجود ندارد")
                for row in ws.iter_rows(min_row=2,values_only=True):
                    ip=text(row[mp["ip_address"]]).strip();
                    if not ip:continue
                    ipaddress.ip_address(ip);x={k:text(row[i]) if i is not None and i<len(row) else "" for k,i in mp.items() if k!="ip_address"};x["ip_address"]=ip;ts=datetime.now().strftime("%Y-%m-%d %H:%M:%S");fields=["ip_address","name","model","vendor","serial_number","ssid","frequency","channel","bandwidth","mode","radio_name","tx_power","rx_power","signal_strength","peer_tx_signal","tx_rate","rx_rate","ccq","snr","noise_floor","notes"];vals=[x.get(f,"") for f in fields];conn.execute(f"INSERT INTO manual_radios({','.join(fields)},created_at,updated_at) VALUES({','.join('?' for _ in fields)},?,?) ON CONFLICT(ip_address) DO UPDATE SET "+','.join(f'{f}=excluded.{f}' for f in fields if f!='ip_address')+",updated_at=excluded.updated_at",vals+[ts,ts]);count+=1
                conn.commit();return {"status":"ok","updated":count}
            # device import - manual/site information only.
            aliases={"ip":"ip","ip address":"ip","latitude":"latitude","longitude":"longitude","province":"province","city":"city","antenna gain":"antenna_gain","gain":"antenna_gain","polarization":"polarization","capacity":"capacity"};mp={aliases.get(h,h):i for h,i in pos.items()}
            if "ip" not in mp:raise ValueError("ستون IP در Excel وجود ندارد")
            count=0;not_found=[]
            for row in ws.iter_rows(min_row=2,values_only=True):
                ip=text(row[mp["ip"]]).strip();
                if not ip:continue
                ip=str(ipaddress.ip_address(ip));found=conn.execute("SELECT 1 FROM devices WHERE ip_address=?",(ip,)).fetchone()
                if not found:not_found.append(ip);continue
                mapping={"latitude":"manual_latitude","longitude":"manual_longitude","province":"manual_province","city":"manual_city","antenna_gain":"manual_antenna_gain","polarization":"manual_polarization","capacity":"manual_capacity"};sets=[];args=[]
                for eff,man in mapping.items():
                    if eff in mp and row[mp[eff]] not in (None,""):val=text(row[mp[eff]]);sets += [f"{man}=?",f"{eff}=?"];args += [val,val]
                if sets:sets.append("last_updated=?");args += [datetime.now().strftime("%Y-%m-%d %H:%M:%S"),ip];conn.execute(f"UPDATE devices SET {','.join(sets)} WHERE ip_address=?",args);count+=1
            conn.commit();return {"status":"ok","updated":count,"not_found":len(not_found),"not_found_ips":not_found}
        except Exception:conn.rollback();raise
        finally:conn.close()
    finally:wb.close()

def template_excel(mode="devices"):
    wb=Workbook();ws=wb.active;ws.title=mode.title()
    if mode=="blacklist":headers=BLACK_HEADERS;row=["172.17.250.0/24","نمونه",1]
    elif mode=="radios":headers=RADIO_HEADERS;row=["172.17.240.118","Radio-1","NetMetal","MikroTik","","SSID","5855","20","20 MHz","station-bridge","APKavosifar5195","20 dBm","-47 dBm","-47 dBm","-29 dBm","130 Mbps","175.5 Mbps","32%","43 dB","-90 dBm","یادداشت"]
    else:headers=["IP","Latitude","Longitude","Province","City","Antenna Gain","Polarization","Capacity"];row=["172.17.240.118","35.70","51.40","تهران","تهران","30 dBi","H","500 Mbps"]
    ws.append(headers);ws.append(row);sheet_style(ws);out=io.BytesIO();wb.save(out);wb.close();out.seek(0);return out
